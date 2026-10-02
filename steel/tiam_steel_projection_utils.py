"""Replay the final steel workbook's regional production model.

Country demand and scrap are kt; population is thousands of people. Production
uses normalized 2019 calibration shares, a constant 2019-2050 per-capita demand
CAGR (with explicit regional adjustments), and China as the balancing region.
"""

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


WORKBOOK_PATH = Path(__file__).resolve().parent / "outputs" / "Steel_demand_and_scrap_projections [SP].xlsx"
BASE_YEAR = 2019
CAGR_END_YEAR = 2050
END_YEAR = 2060
YEARS = [str(year) for year in range(BASE_YEAR, END_YEAR + 1)]
COUNTRY_METRICS = {"EUSC": "demand", "POPN": "population", "SCRAP": "scrap"}
REGIONAL_TABLE_ROWS = {
    "demand": 3,
    "population": 36,
    "per_capita": 68,
    "production": 100,
    "scrap": 168,
}


def _numeric(frame: pd.DataFrame, columns: list[str], label: str) -> pd.DataFrame:
    missing = set(columns) - set(frame.columns)
    if missing:
        raise ValueError(f"{label} is missing columns: {sorted(missing)}")
    values = frame[columns].apply(pd.to_numeric, errors="raise")
    if not np.isfinite(values.to_numpy(dtype=float)).all():
        raise ValueError(f"{label} contains missing or nonfinite values")
    if values.lt(0).any().any():
        raise ValueError(f"{label} contains negative values")
    return values


def _require_formula(sheet, row: int, column: int, expected: str) -> None:
    cell = sheet.cell(row, column)
    if cell.value != expected:
        raise ValueError(f"Unexpected workbook formula at {sheet.title}!{cell.coordinate}: {cell.value!r}; expected {expected!r}")


def _validate_source_formulas(workbook, country_rows: list[int]) -> None:
    """Reject source changes that would invalidate the replayed methodology."""
    data = workbook["Data"]
    regional = workbook["OMNIA_Data"]
    for row in country_rows:
        _require_formula(data, row, 3, f"=VLOOKUP(B{row},CountryCodes!$B$4:$C$179,2,FALSE)")
    for offset in range(28):
        production_row = 100 + offset
        per_capita_row = 68 + offset
        adjustment = regional.cell(production_row, 57).value
        _require_formula(regional, production_row, 51, f"=AX{production_row}/SUM($AX$100:$AX$127)")
        _require_formula(regional, production_row, 54, f"=VLOOKUP(BA{production_row},$AW$100:$AY$127,3,FALSE)")
        growth_formula = f"=AV{per_capita_row}"
        if adjustment is not None:
            growth_formula += f"*BE{production_row}"
        _require_formula(regional, production_row, 56, growth_formula)
        _require_formula(regional, per_capita_row, 48, f"=(AJ{per_capita_row}/E{per_capita_row})^(1/(AJ$2-E$2))-1")
        for column in range(5, 47):
            letter = get_column_letter(column)
            previous = get_column_letter(column - 1)
            source_letter = get_column_letter(column + 71)
            _require_formula(regional, 3 + offset, column, f"={letter}{36 + offset}*{letter}{68 + offset}/1000")
            _require_formula(regional, 36 + offset, column, f"=Data!{source_letter}{754 + offset}")
            _require_formula(regional, per_capita_row, column, f"=Data!{source_letter}{782 + offset}")
            _require_formula(regional, 168 + offset, column, f"=Data!{source_letter}{810 + offset}")
            _require_formula(data, 782 + offset, column + 71, f"={source_letter}{726 + offset}/{source_letter}{754 + offset}*1000")
            for aggregate_row in (726 + offset, 754 + offset, 810 + offset):
                _require_formula(data, aggregate_row, column + 71, f"=SUMIFS({source_letter}$4:{source_letter}$720,$F$4:$F$720,$F{aggregate_row},$C$4:$C$720,$C{aggregate_row})")
            if production_row == 108:
                expected = f"={letter}130-SUM({letter}109:{letter}127,{letter}100:{letter}107)"
            elif column == 5:
                expected = f"=BB{production_row}*E$32"
            else:
                expected = f"={previous}{production_row}*(1+$BD{production_row})"
            _require_formula(regional, production_row, column, expected)
    for column in range(5, 47):
        letter = get_column_letter(column)
        _require_formula(regional, 32, column, f"=SUM({letter}3:{letter}30)/1000")
        _require_formula(regional, 130, column, f"={letter}32")


def read_workbook_inputs(path: Path = WORKBOOK_PATH) -> dict:
    """Read and validate saved country inputs and the workbook model structure.

    Returned calibration values are the raw model calibration, converted from
    Mt to kt. They differ from final 2019 anchors, which normalize these shares
    to the workbook's global 2019 demand. No country calibration is inferred.
    """
    path = Path(path)
    cached_workbook = load_workbook(path, data_only=True)
    formula_workbook = load_workbook(path, data_only=False)
    try:
        required_sheets = {"Data", "CountryCodes", "OMNIA_Data"}
        if not required_sheets.issubset(cached_workbook.sheetnames):
            raise ValueError(f"Steel workbook needs sheets {sorted(required_sheets)}")
        data = cached_workbook["Data"]
        country_codes = cached_workbook["CountryCodes"]
        regional = cached_workbook["OMNIA_Data"]
        year_columns = {str(data.cell(3, column).value): column for column in range(8, data.max_column + 1)}
        if any(year not in year_columns for year in YEARS):
            raise ValueError("Country workbook headers must cover 2019-2060")
        for column, year in enumerate(range(BASE_YEAR, END_YEAR + 1), start=5):
            if regional.cell(2, column).value != year:
                raise ValueError("Regional workbook headers must cover 2019-2060")
            if year_columns[str(year)] != column + 71:
                raise ValueError("Country workbook year columns no longer match source formulas")

        countries = {metric: [] for metric in COUNTRY_METRICS.values()}
        source_rows = []
        for row in range(4, 721):
            iso3 = data.cell(row, 2).value
            metric = COUNTRY_METRICS.get(data.cell(row, 6).value)
            if not isinstance(iso3, str) or len(iso3) != 3 or not iso3.isupper() or metric is None:
                continue
            countries[metric].append({
                "Country": data.cell(row, 1).value,
                "ISO3": iso3,
                "OMNIARegion": data.cell(row, 3).value,
                **{year: data.cell(row, year_columns[year]).value for year in YEARS},
            })
            source_rows.append(row)
        countries = {metric: pd.DataFrame(rows) for metric, rows in countries.items()}
        for metric, frame in countries.items():
            if len(frame) != 176 or frame["ISO3"].duplicated().any():
                raise ValueError(f"Expected 176 unique country {metric} series")
            frame[YEARS] = _numeric(frame, YEARS, f"Country {metric}")
        country_sets = [set(frame["ISO3"]) for frame in countries.values()]
        if any(codes != country_sets[0] for codes in country_sets):
            raise ValueError("Country demand, population and scrap coverage differs")

        mapping = pd.DataFrame([
            {"Country": country_codes.cell(row, 1).value,
             "ISO3": country_codes.cell(row, 2).value,
             "OMNIARegion": country_codes.cell(row, 3).value}
            for row in range(4, 180)
        ])
        if mapping["ISO3"].duplicated().any() or set(mapping["ISO3"]) != country_sets[0]:
            raise ValueError("Workbook CountryCodes coverage differs from country inputs")
        lookup = mapping.set_index("ISO3")["OMNIARegion"]
        for metric, frame in countries.items():
            if frame["OMNIARegion"].isna().any() or not frame["ISO3"].map(lookup).eq(frame["OMNIARegion"]).all():
                raise ValueError(f"Country {metric} cached regions disagree with CountryCodes")

        calibration = pd.DataFrame([
            {"OMNIARegion": regional.cell(row, 49).value,
             "Calibration2019_kt": regional.cell(row, 50).value * 1000.0}
            for row in range(100, 128)
        ])
        adjustments = pd.DataFrame([
            {"OMNIARegion": regional.cell(row, 2).value,
             "AdjustmentFactor": regional.cell(row, 57).value if regional.cell(row, 57).value is not None else 1.0}
            for row in range(100, 128)
        ])
        cached = {}
        for metric, first_row in REGIONAL_TABLE_ROWS.items():
            rows = []
            for row in range(first_row, first_row + 28):
                record = {"OMNIARegion": regional.cell(row, 2).value}
                for column, year in enumerate(YEARS, start=5):
                    value = regional.cell(row, column).value
                    record[year] = value * 1000.0 if metric == "production" else value
                rows.append(record)
            cached[metric] = pd.DataFrame(rows)
            if cached[metric]["OMNIARegion"].duplicated().any():
                raise ValueError(f"Duplicate regions in workbook {metric} table")
            cached[metric][YEARS] = _numeric(cached[metric], YEARS, f"Workbook {metric}")
        region_sets = [set(frame["OMNIARegion"]) for frame in cached.values()]
        if any(regions != region_sets[0] for regions in region_sets):
            raise ValueError("Workbook regional table coverage differs")
        if set(mapping["OMNIARegion"]) != region_sets[0] or len(region_sets[0]) != 28:
            raise ValueError("Expected 28 consistently mapped OMNIA workbook regions")
        if regional.cell(108, 2).value != "CHN":
            raise ValueError("Workbook balancing production region is no longer CHN")
        _validate_source_formulas(formula_workbook, source_rows)
        return {"countries": countries, "mapping": mapping,
                "calibration": calibration, "adjustments": adjustments,
                "cached": cached, "source_path": path}
    finally:
        cached_workbook.close()
        formula_workbook.close()


def calculate_regional_model(
    countries: dict[str, pd.DataFrame],
    mapping: pd.DataFrame,
    calibration: pd.DataFrame,
    adjustments: pd.DataFrame,
    region_column: str = "TIAMRegion",
    residual_region: str = "CHI",
    end_year: int = END_YEAR,
) -> dict[str, pd.DataFrame]:
    """Apply the workbook formulas using explicitly supplied regional inputs."""
    if end_year < CAGR_END_YEAR:
        raise ValueError("The model horizon must include its 2050 CAGR endpoint")
    years = [str(year) for year in range(BASE_YEAR, end_year + 1)]
    if not {"ISO3", region_column}.issubset(mapping.columns):
        raise ValueError(f"Mapping needs ISO3 and {region_column}")
    mapping = mapping[["ISO3", region_column]].copy()
    for column in mapping:
        mapping[column] = mapping[column].fillna("").astype(str).str.strip()
    if not mapping["ISO3"].str.fullmatch("[A-Z]{3}").all() or mapping[region_column].eq("").any():
        raise ValueError("Regional mapping contains invalid or blank country/region codes")
    if mapping["ISO3"].duplicated().any():
        raise ValueError("Regional mapping contains duplicate ISO3 assignments")

    parameters = calibration[[region_column, "Calibration2019_kt"]].copy()
    if parameters.empty or parameters[region_column].isna().any() or parameters[region_column].duplicated().any():
        raise ValueError("Calibration needs unique nonblank regional rows")
    parameters["Calibration2019_kt"] = _numeric(parameters, ["Calibration2019_kt"], "2019 calibration")
    if parameters["Calibration2019_kt"].sum() <= 0:
        raise ValueError("Global raw 2019 production calibration must be positive")
    if residual_region not in set(parameters[region_column]):
        raise ValueError(f"Missing balancing region {residual_region}")
    if not {region_column, "AdjustmentFactor"}.issubset(adjustments.columns):
        raise ValueError("Adjustments need one explicit factor for every model region")
    if adjustments[region_column].duplicated().any() or set(adjustments[region_column]) != set(parameters[region_column]):
        raise ValueError("Adjustment regions differ from calibration regions")
    parameters = parameters.merge(adjustments[[region_column, "AdjustmentFactor"]], on=region_column, validate="one_to_one", sort=False)
    parameters["AdjustmentFactor"] = _numeric(parameters, ["AdjustmentFactor"], "Production adjustments")
    if parameters["AdjustmentFactor"].le(0).any():
        raise ValueError("Production adjustment factors must be positive")
    regions = parameters[region_column].tolist()
    lookup = mapping.set_index("ISO3")[region_column]
    result = {}
    country_set = None
    for metric in ("demand", "population", "scrap"):
        if metric not in countries:
            raise ValueError(f"Missing country {metric} inputs")
        source = countries[metric].copy()
        source.columns = source.columns.map(str)
        if "ISO3" not in source or source.empty or source["ISO3"].duplicated().any():
            raise ValueError(f"Country {metric} needs unique ISO3 series")
        current_set = set(source["ISO3"])
        if country_set is not None and current_set != country_set:
            raise ValueError("Country demand, population and scrap coverage differs")
        country_set = current_set
        assigned = source["ISO3"].map(lookup)
        if assigned.isna().any():
            raise ValueError(f"Unmapped country {metric} codes: {sorted(source.loc[assigned.isna(), 'ISO3'])}")
        if set(assigned) - set(regions):
            raise ValueError("Country mapping includes regions absent from calibration")
        values = _numeric(source, years, f"Country {metric}")
        totals = values.groupby(assigned.rename(region_column)).sum().reindex(regions, fill_value=0.0)
        if not np.allclose(totals.sum(), values.sum(), rtol=1e-12, atol=1e-6):
            raise ValueError(f"Regional aggregation loses country {metric} volume")
        result[metric] = totals

    if result["population"].le(0).any().any():
        raise ValueError("Every model region needs positive annual population")
    per_capita = result["demand"] / result["population"] * 1000.0
    if per_capita[[str(BASE_YEAR), str(CAGR_END_YEAR)]].le(0).any().any():
        raise ValueError("Every model region needs positive 2019 and 2050 per-capita demand")
    cagr = (per_capita[str(CAGR_END_YEAR)] / per_capita[str(BASE_YEAR)]) ** (1.0 / (CAGR_END_YEAR - BASE_YEAR)) - 1.0
    parameters = parameters.set_index(region_column)
    parameters["ProductionShare2019"] = parameters["Calibration2019_kt"] / parameters["Calibration2019_kt"].sum()
    parameters["PerCapita2019_kg"] = per_capita[str(BASE_YEAR)]
    parameters["PerCapita2050_kg"] = per_capita[str(CAGR_END_YEAR)]
    parameters["PerCapitaCAGR"] = cagr
    parameters["AppliedProductionCAGR"] = cagr * parameters["AdjustmentFactor"]
    production = pd.DataFrame(index=pd.Index(regions, name=region_column), columns=years, dtype=float)
    production[str(BASE_YEAR)] = parameters["ProductionShare2019"] * result["demand"][str(BASE_YEAR)].sum()
    non_residual = production.index != residual_region
    for year in range(BASE_YEAR + 1, end_year + 1):
        production.loc[non_residual, str(year)] = production.loc[non_residual, str(year - 1)] * (1.0 + parameters.loc[non_residual, "AppliedProductionCAGR"])
    for year in years:
        production.loc[residual_region, year] = result["demand"][year].sum() - production.loc[non_residual, year].sum()
    _numeric(production, years, "Calculated production (including China residual)")
    if not np.allclose(production.sum(), result["demand"].sum(), rtol=1e-12, atol=1e-6):
        raise ValueError("Calculated production does not balance global steel demand")
    result["per_capita"] = per_capita
    result["production"] = production
    result = {metric: frame.reset_index() for metric, frame in result.items()}
    result["parameters"] = parameters.reset_index()
    return result


def replay_omnia_model(inputs: dict) -> dict[str, pd.DataFrame]:
    """Validate a complete 2019-2060 replay against the workbook's saved values."""
    result = calculate_regional_model(
        inputs["countries"], inputs["mapping"], inputs["calibration"], inputs["adjustments"],
        region_column="OMNIARegion", residual_region="CHN",
    )
    for metric, cached in inputs["cached"].items():
        expected = cached.set_index("OMNIARegion")[YEARS]
        actual = result[metric].set_index("OMNIARegion").reindex(expected.index)[YEARS]
        if not np.allclose(actual, expected, rtol=1e-12, atol=1e-6):
            difference = np.abs(actual.to_numpy(dtype=float) - expected.to_numpy(dtype=float)).max()
            raise ValueError(f"Workbook {metric} method replay failed; maximum difference {difference:.12g}")
    return result
