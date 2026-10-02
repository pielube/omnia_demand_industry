"""Rebuild the steel workbook's calculation for TIAM and export its CSVs."""

import argparse
from pathlib import Path
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from tiam_projection_utils import YEARS, calculate_growth_rates, read_tiam_mapping
from steel import rebase_steel_production_omnia_2019_worldsteel_indexed as wsa
from steel.tiam_steel_projection_utils import (
    calculate_regional_model,
    read_workbook_inputs,
    replay_omnia_model,
)

OUTPUT_DIR = ROOT / "steel" / "outputs"
SOURCE_WORKBOOK = OUTPUT_DIR / "Steel_demand_and_scrap_projections [SP].xlsx"
TIAM_WORKBOOK = OUTPUT_DIR / "Steel_demand_and_scrap_projections [TIAM].xlsx"
MODEL_YEARS = [str(year) for year in range(2019, 2061)]
DESCRIPTIONS = {
    "AFR": "Africa", "AUS": "Australia and New Zealand", "CAN": "Canada",
    "CHI": "China", "CSA": "Central and South America", "EEU": "Eastern Europe",
    "FSU": "Former Soviet Union", "IND": "India", "JPN": "Japan",
    "MEA": "Middle East", "MEX": "Mexico", "ODA": "Other developing Asia",
    "SKO": "South Korea", "UK": "United Kingdom", "USA": "United States",
    "WEU": "Western Europe",
}


def build_calibration(inputs, mapping, method):
    """Allocate the source calibration, making its new assumptions explicit."""
    historical, _ = wsa.read_historical_source()
    source_map = wsa.read_source_map(
        wsa.HISTORICAL_MAP_CSV, set(historical["SourceCountry"]), "World Steel 2019-2020"
    )
    if wsa.file_sha256(wsa.HISTORICAL_MAP_CSV) != wsa.EXPECTED_HISTORICAL_MAP_SHA256:
        raise ValueError("Historical World Steel country map has changed")
    wsa.validate_map_against_shared(source_map, wsa.read_shared_map(), "2019-2020 source")
    country = source_map.merge(historical, on="SourceCountry", validate="one_to_one")
    country = country.loc[country["Availability2019"].eq("reported")].copy()
    country = country.merge(mapping[["ISO3", "TIAMRegion"]], on="ISO3", how="left", validate="one_to_one")
    if country["TIAMRegion"].isna().any():
        raise ValueError("World Steel calibration countries missing from TIAM mapping")
    regional = inputs["calibration"].set_index("OMNIARegion")["Calibration2019_kt"]
    adjustments = inputs["adjustments"].set_index("OMNIARegion")["AdjustmentFactor"]
    country["SourceAdjustmentFactor"] = country["OMNIARegion"].map(adjustments)
    if country["SourceAdjustmentFactor"].isna().any():
        raise ValueError("World Steel countries missing source growth adjustments")
    observed_totals = country.groupby("OMNIARegion")["Production2019_kt"].sum()
    missing = set(regional.index) - set(observed_totals.loc[observed_totals.gt(0)].index)
    if missing:
        raise ValueError(f"Source calibration regions lack positive 2019 country weights: {sorted(missing)}")
    country["SourceRegionShare2019"] = (
        country["Production2019_kt"] / country["OMNIARegion"].map(observed_totals)
    )
    if method == "allocated-model":
        country["Calibration2019_kt"] = (
            country["SourceRegionShare2019"] * country["OMNIARegion"].map(regional)
        )
        actual = country.groupby("OMNIARegion")["Calibration2019_kt"].sum().reindex(regional.index)
        if not np.allclose(actual, regional, rtol=1e-13, atol=1e-8):
            raise ValueError("Country allocation does not preserve source model calibration")
    elif method == "worldsteel":
        country["Calibration2019_kt"] = country["Production2019_kt"]
    else:
        raise ValueError(f"Unknown calibration method: {method}")
    country["AdjustmentWeightedCalibration"] = (
        country["Calibration2019_kt"] * country["SourceAdjustmentFactor"]
    )
    totals = country.groupby("TIAMRegion", as_index=False).agg(
        Calibration2019_kt=("Calibration2019_kt", "sum"),
        AdjustmentWeightedCalibration=("AdjustmentWeightedCalibration", "sum"),
    )
    totals["AdjustmentFactor"] = totals["AdjustmentWeightedCalibration"] / totals["Calibration2019_kt"]
    if set(totals["TIAMRegion"]) != set(DESCRIPTIONS):
        raise ValueError("Not every TIAM region has a production calibration")
    country["CalibrationMethod"] = method
    country["CalibrationSource"] = "OMNIA_Data!AW100:AX127" if method == "allocated-model" else "World Steel 2019"
    country["AdjustmentMethod"] = "Calibration-weighted mean of source OMNIA factors"
    country["SourceWorkbookSHA256"] = wsa.file_sha256(SOURCE_WORKBOOK)
    country["HistoricalSourceCSV_SHA256"] = wsa.file_sha256(wsa.HISTORICAL_SOURCE_CSV)
    return (
        totals[["TIAMRegion", "Calibration2019_kt"]],
        totals[["TIAMRegion", "AdjustmentFactor"]],
        country.drop(columns="AdjustmentWeightedCalibration").sort_values(["TIAMRegion", "ISO3"]),
    )


def clear_cells(sheet, first_row, last_row, first_column=1, last_column=None):
    last_column = last_column or sheet.max_column
    for row in sheet.iter_rows(min_row=first_row, max_row=last_row, min_col=first_column, max_col=last_column):
        for cell in row:
            cell.value = None


def build_workbook(mapping, calibration, adjustments, allocation, path=TIAM_WORKBOOK):
    """Keep the source layout and country history, replacing its regional formulas."""
    workbook = load_workbook(SOURCE_WORKBOOK, data_only=False)
    regional_sheet = workbook["OMNIA_Data"]
    regional_sheet.title = "TIAM_Data"
    data = workbook["Data"]
    codes = workbook["CountryCodes"]
    lookup = mapping.set_index("ISO3")["TIAMRegion"]
    regions = sorted(DESCRIPTIONS)
    map_sheet = workbook.create_sheet("TIAM_Mapping")
    map_sheet.append(["Country", "ISO3", "TIAMRegion"])
    for row in mapping[["Country", "ISO3", "TIAMRegion"]].itertuples(index=False, name=None):
        map_sheet.append(list(row))
    codes["C3"] = "TIAM_Region"
    clear_cells(codes, 2, codes.max_row, 7, 20)
    codes["G2"] = "Source: shared_inputs/tiam_country_countrycode_region.csv and supplements"
    for column, value in zip((7, 9, 11), ("Country", "ISO3", "TIAMRegion")):
        codes.cell(3, column, value)
    for number, row in enumerate(mapping[["Country", "ISO3", "TIAMRegion"]].itertuples(index=False, name=None), 4):
        for column, value in zip((7, 9, 11), row):
            codes.cell(number, column, value)
    for row in range(4, 180):
        iso3 = codes.cell(row, 2).value
        if iso3 not in lookup.index:
            raise ValueError(f"Workbook country missing TIAM mapping: {iso3}")
        codes.cell(row, 3, f'=VLOOKUP(B{row},TIAM_Mapping!$B$2:$C${len(mapping)+1},2,FALSE)')
    data["C3"] = "TIAM_Region"
    data["C725"] = "TIAM_Region"
    for row in range(4, 721):
        if data.cell(row, 2).value == "Global":
            data.cell(row, 3).value = None
        elif data.cell(row, 2).value:
            data.cell(row, 3, f'=VLOOKUP(B{row},CountryCodes!$B$4:$C$179,2,FALSE)')
    clear_cells(data, 726, 837)
    blocks = [(726, "EUSC", "End-use steel consumption [kt]", "kt"),
              (754, "POPN", "Population", "000s"),
              (782, None, "End-use steel consumption per capita [kg/person]", "kg / cap"),
              (810, "SCRAP", "Total available scrap", "kt")]
    for start, metric_code, metric, unit in blocks:
        for offset, region in enumerate(regions):
            row = start + offset
            for column, value in ((3, region), (5, metric), (6, metric_code), (7, unit)):
                data.cell(row, column).value = value
            for column in range(8, 158):
                letter = get_column_letter(column)
                if metric_code is None:
                    formula = f'={letter}{726+offset}/{letter}{754+offset}*1000'
                else:
                    formula = f'=SUMIFS({letter}$4:{letter}$720,$F$4:$F$720,$F{row},$C$4:$C$720,$C{row})'
                data.cell(row, column, formula)
    starts = (3, 36, 68, 100, 134, 168, 201)
    for start in starts:
        clear_cells(regional_sheet, start, start + 27)
    # Remove the original OMNIA calibration/adjustment table before rebuilding.
    clear_cells(regional_sheet, 99, 127, 49, 93)
    labels = {3: "Consumption (kt)", 36: "Population (000s)", 68: "Consumption (kg/person)",
              100: "Production (Mt)", 134: "Production (index)", 168: "Scrap (kt)", 201: "Scrap (index)"}
    raw = calibration.set_index("TIAMRegion")["Calibration2019_kt"]
    factors = adjustments.set_index("TIAMRegion")["AdjustmentFactor"]
    for coordinate, label in {"AW99": "TIAM region", "AX99": "2019 calibration (Mt)",
                              "AY99": "Production share (2019)", "BA99": "TIAM region",
                              "BB99": "Production share (2019)", "BD99": "Production growth rate",
                              "BE99": "Adjustment"}.items():
        regional_sheet[coordinate] = label
    residual = 100 + regions.index("CHI")
    for offset, region in enumerate(regions):
        for start in starts:
            row = start + offset
            for column, value in ((2, region), (3, DESCRIPTIONS[region]), (4, labels[start])):
                regional_sheet.cell(row, column, value)
        row = 100 + offset
        for column, value in ((49, region), (50, float(raw[region])/1000), (53, region), (57, float(factors[region]))):
            regional_sheet.cell(row, column, value)
        regional_sheet[f"AY{row}"] = f'=AX{row}/SUM($AX$100:$AX$115)'
        regional_sheet[f"BB{row}"] = f'=VLOOKUP(BA{row},$AW$100:$AY$115,3,FALSE)'
        regional_sheet[f"BD{row}"] = f'=AV{68+offset}*BE{row}'
        for column in range(5, 47):
            letter = get_column_letter(column)
            prior = get_column_letter(column - 1)
            source_column = get_column_letter(column + 71)
            regional_sheet.cell(3+offset, column, f'={letter}{36+offset}*{letter}{68+offset}/1000')
            regional_sheet.cell(36+offset, column, f'=Data!{source_column}{754+offset}')
            regional_sheet.cell(68+offset, column, f'=Data!{source_column}{782+offset}')
            if region == "CHI":
                production = f'={letter}130-SUM({letter}100:{letter}{residual-1},{letter}{residual+1}:{letter}115)'
            elif column == 5:
                production = f'=BB{row}*E$32'
            else:
                production = f'={prior}{row}*(1+$BD{row})'
            regional_sheet.cell(row, column, production)
            regional_sheet.cell(134+offset, column, f'={letter}{row}/$E{row}')
            regional_sheet.cell(168+offset, column, f'=Data!{source_column}{810+offset}')
            regional_sheet.cell(201+offset, column, f'={letter}{168+offset}/$E{168+offset}')
        for start in (3, 68):
            cagr_row = start + offset
            regional_sheet[f"AV{cagr_row}"] = f'=(AJ{cagr_row}/E{cagr_row})^(1/(AJ$2-E$2))-1'
    for column in range(5, 47):
        letter = get_column_letter(column)
        for row, formula in ((32, f'=SUM({letter}3:{letter}18)/1000'),
                             (129, f'=SUM({letter}100:{letter}115)'), (130, f'={letter}32'),
                             (197, f'=SUM({letter}168:{letter}183)/1000'),
                             (230, f'={letter}197/$E197')):
            regional_sheet.cell(row, column, formula)
    audit = workbook.create_sheet("Production_Calibration")
    audit.append(list(allocation.columns))
    for row in allocation.itertuples(index=False, name=None):
        audit.append(list(row))
    notes = workbook.create_sheet("Method")
    notes.append(["Source workbook", SOURCE_WORKBOOK.name])
    notes.append(["Source SHA256", wsa.file_sha256(SOURCE_WORKBOOK)])
    notes.append(["Method", "2019 calibration shares times global 2019 demand; constant adjusted per-capita CAGR; China residual"])
    notes.append(["Allocation assumption", allocation["CalibrationMethod"].iloc[0]])
    notes.append(["Adjustment assumption", "2019 calibration-weighted mean of the source OMNIA regional factors"])
    notes.append(["Units", "Demand/scrap: kt; population: thousands; per-capita demand: kg/person; production: Mt"])
    notes.append(["CSV horizon", "2019-2050; milestone indices retain 2050 values through 2100"])
    workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    workbook.active = workbook.sheetnames.index("TIAM_Data")
    workbook.save(path)
    workbook.close()


def recalculate_workbook(path):
    """Use a separate hidden Excel instance to save recalculated formula caches."""
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    application = None
    workbook = None
    try:
        application = win32com.client.DispatchEx("Excel.Application")
        application.Visible = False
        application.DisplayAlerts = False
        application.AskToUpdateLinks = False
        application.AutomationSecurity = 3
        workbook = application.Workbooks.Open(str(Path(path).resolve()), UpdateLinks=0, ReadOnly=False)
        application.CalculateFullRebuild()
        workbook.Save()
        workbook.Close(False)
        workbook = None
    finally:
        if workbook is not None:
            workbook.Close(False)
        if application is not None:
            application.Quit()
        pythoncom.CoUninitialize()


def validate_workbook(model, path=TIAM_WORKBOOK):
    workbook = load_workbook(path, data_only=True)
    sheet = workbook["TIAM_Data"]
    errors = [f"{tab.title}!{cell.coordinate}: {cell.value}" for tab in workbook
              for row in tab for cell in row if cell.data_type == "e"]
    if errors:
        workbook.close()
        raise ValueError(f"Recalculated workbook contains Excel errors: {errors[:10]}")
    extracted = {}
    for name, start, scale in (("demand", 3, 1), ("population", 36, 1), ("per_capita", 68, 1),
                               ("production", 100, 1000), ("scrap", 168, 1)):
        expected = model[name].set_index("TIAMRegion")
        records = []
        for offset, region in enumerate(sorted(DESCRIPTIONS)):
            if sheet.cell(start+offset, 2).value != region:
                raise ValueError(f"Unexpected {name} region in recalculated TIAM workbook")
            actual = np.array([sheet.cell(start+offset, column).value for column in range(5, 47)], dtype=float)*scale
            if not np.isfinite(actual).all() or not np.allclose(actual, expected.loc[region, MODEL_YEARS].to_numpy(dtype=float), rtol=1e-12, atol=1e-6):
                raise ValueError(f"Recalculated workbook does not match {name} model for {region}")
            records.append({"TIAMRegion": region, **dict(zip(MODEL_YEARS, actual))})
        extracted[name] = pd.DataFrame(records)
    for metric, start in (("production", 134), ("scrap", 201)):
        expected = extracted[metric].set_index("TIAMRegion")
        for offset, region in enumerate(sorted(DESCRIPTIONS)):
            actual = np.array([sheet.cell(start+offset, column).value for column in range(5, 47)], dtype=float)
            index = expected.loc[region, MODEL_YEARS].to_numpy(dtype=float) / float(expected.at[region, "2019"])
            if sheet.cell(start+offset, 2).value != region or not np.allclose(actual, index, rtol=1e-13, atol=1e-13):
                raise ValueError(f"Workbook {metric} growth indices do not match {region}")
    workbook.close()
    return extracted


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--calibration-method", choices=("allocated-model", "worldsteel"), required=True)
    args = parser.parse_args()
    inputs = read_workbook_inputs(SOURCE_WORKBOOK)
    replay_omnia_model(inputs)
    mapping = read_tiam_mapping()
    calibration, adjustments, allocation = build_calibration(inputs, mapping, args.calibration_method)
    model = calculate_regional_model(inputs["countries"], mapping, calibration, adjustments)
    from steel.rebase_steel_production_tiam_2019_worldsteel_indexed import build_outputs as build_indexed
    build_workbook(mapping, calibration, adjustments, allocation)
    recalculate_workbook(TIAM_WORKBOOK)
    extracted = validate_workbook(model)
    csv_outputs = {}
    for metric in ("production", "scrap"):
        projection = extracted[metric][["TIAMRegion", *YEARS]]
        csv_outputs[OUTPUT_DIR / f"steel_{metric}_tiam.csv"] = projection
        csv_outputs[OUTPUT_DIR / f"steel_{metric}_tiam_growth_rates.csv"] = calculate_growth_rates(projection)
    csv_outputs.update(build_indexed(csv_outputs[OUTPUT_DIR / "steel_production_tiam.csv"], mapping))
    for path, frame in csv_outputs.items():
        frame.to_csv(path, index=False, float_format="%.17g")
        print(f"Saved: {path.relative_to(ROOT)}")
    print(f"Saved: {TIAM_WORKBOOK.relative_to(ROOT)}")
    print("Validated original OMNIA method, 16 TIAM regions, Excel recalculation, and World Steel indexing")


if __name__ == "__main__":
    main()
