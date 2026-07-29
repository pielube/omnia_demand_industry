"""Extract OMNIA steel production and scrap projections from the final workbook."""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string


BASE_DIR = Path(__file__).resolve().parent
OUTPUTS_DIR = BASE_DIR / "outputs"
WORKBOOK_PATH = OUTPUTS_DIR / "Steel_demand_and_scrap_projections [SP].xlsx"
SHEET_NAME = "OMNIA_Data"

BASE_YEAR = 2019
END_YEAR = 2050
RANGE_END_COLUMN = column_index_from_string("AT")
MILESTONE_YEARS = [
    2019,
    2023,
    2025,
    2030,
    2035,
    2040,
    2045,
    2050,
    2060,
    2070,
    2080,
    2090,
    2100,
]
SOURCE_MILESTONE_YEARS = [year for year in MILESTONE_YEARS if year <= END_YEAR]

TABLES = {
    "production": {
        "header_row": 99,
        "first_data_row": 100,
        "last_data_row": 127,
        "global_row": 129,
        "metric": "Production (Kt)",
        # Regional production values are numerically Mt despite the row label;
        # their sum matches the workbook's global production value in Mt.
        "output_scale": 1000.0,
        "global_scale": 1000.0,
        "projection_csv": OUTPUTS_DIR / "steel_production_omnia.csv",
        "growth_csv": OUTPUTS_DIR / "steel_production_omnia_growth_rates.csv",
    },
    "scrap": {
        "header_row": 167,
        "first_data_row": 168,
        "last_data_row": 195,
        "global_row": 197,
        "metric": "Scrap (kt)",
        "output_scale": 1.0,
        # Regional scrap values are kt and the global workbook row is Mt.
        "global_scale": 1000.0,
        "projection_csv": OUTPUTS_DIR / "steel_scrap_omnia.csv",
        "growth_csv": OUTPUTS_DIR / "steel_scrap_omnia_growth_rates.csv",
    },
}


def extract_table(worksheet, name: str, config: dict) -> pd.DataFrame:
    """Extract one specified B:AT table and standardise output values to kt."""
    header = [
        worksheet.cell(config["header_row"], column).value
        for column in range(2, RANGE_END_COLUMN + 1)
    ]
    year_columns = {}
    for offset, value in enumerate(header):
        if isinstance(value, (int, float)) and float(value).is_integer():
            year = int(value)
            if BASE_YEAR <= year <= END_YEAR:
                year_columns[year] = offset + 2

    expected_years = set(range(BASE_YEAR, END_YEAR + 1))
    missing_years = sorted(expected_years - set(year_columns))
    if missing_years:
        raise ValueError(f"{name} table is missing years: {missing_years}")

    rows = []
    for row_number in range(config["first_data_row"], config["last_data_row"] + 1):
        region = worksheet.cell(row_number, 2).value
        metric = worksheet.cell(row_number, 4).value
        if not region:
            raise ValueError(f"Blank OMNIA region in {name} row {row_number}")
        if metric != config["metric"]:
            raise ValueError(
                f"Unexpected metric in {name} row {row_number}: {metric!r}"
            )

        row = {"OMNIARegion": str(region).strip()}
        for year, column in year_columns.items():
            value = worksheet.cell(row_number, column).value
            if not isinstance(value, (int, float)):
                raise ValueError(
                    f"Non-numeric {name} value in row {row_number}, year {year}: "
                    f"{value!r}"
                )
            if value < 0:
                raise ValueError(
                    f"Negative {name} value in row {row_number}, year {year}"
                )
            row[str(year)] = float(value) * config["output_scale"]
        rows.append(row)

    projection = pd.DataFrame(rows)
    if projection["OMNIARegion"].duplicated().any():
        duplicates = projection.loc[
            projection["OMNIARegion"].duplicated(keep=False), "OMNIARegion"
        ].tolist()
        raise ValueError(f"Duplicate {name} OMNIA regions: {duplicates}")

    # Validate extracted regional totals against the workbook's global row.
    for year, column in year_columns.items():
        global_value = worksheet.cell(config["global_row"], column).value
        if not isinstance(global_value, (int, float)):
            raise ValueError(f"Missing global {name} value for {year}")
        regional_total = projection[str(year)].sum()
        expected_total = float(global_value) * config["global_scale"]
        if abs(regional_total - expected_total) > 1e-5:
            raise ValueError(
                f"{name} regions do not match the global row in {year}: "
                f"{regional_total} vs {expected_total} kt"
            )
    return projection


def calculate_growth_rates(projection: pd.DataFrame) -> pd.DataFrame:
    """Create a milestone-year index with 2019 equal to one."""
    indexed = projection.set_index("OMNIARegion").copy()
    source_years = [str(year) for year in SOURCE_MILESTONE_YEARS]
    base_values = indexed[str(BASE_YEAR)]
    zero_base = base_values.eq(0)
    nonzero_future = indexed.loc[zero_base, source_years].ne(0).any(axis=1)
    if nonzero_future.any():
        regions = indexed.index[zero_base & nonzero_future].tolist()
        raise ValueError(
            "Cannot calculate growth for zero-base regions with "
            f"nonzero future values: {regions}"
        )

    growth = (
        indexed.loc[:, source_years]
        .div(base_values.loc[~zero_base], axis=0)
        .transpose()
    )
    growth.loc[:, zero_base] = 1.0
    growth.index = growth.index.astype(int)

    for year in MILESTONE_YEARS:
        if year > END_YEAR:
            growth.loc[year] = growth.loc[END_YEAR]

    growth = growth.loc[MILESTONE_YEARS]
    growth.index.name = "Year"
    return growth.reset_index()


def main() -> None:
    workbook = load_workbook(WORKBOOK_PATH, read_only=False, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"Workbook has no {SHEET_NAME!r} sheet")
    worksheet = workbook[SHEET_NAME]

    OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)
    for name, config in TABLES.items():
        projection = extract_table(worksheet, name, config)
        growth = calculate_growth_rates(projection)
        projection.to_csv(config["projection_csv"], index=False)
        growth.to_csv(config["growth_csv"], index=False)

        print(f"Saved: {config['projection_csv']}")
        print(f"Saved: {config['growth_csv']}")
        print(f"{name.title()} OMNIA regions: {len(projection)}")
        print(f"Years included: {BASE_YEAR}-{END_YEAR}")


if __name__ == "__main__":
    main()
